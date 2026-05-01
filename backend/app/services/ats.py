import io
import json
import logging
import math
import re
import unicodedata
from collections import Counter
from functools import lru_cache

from fastapi import HTTPException, UploadFile

from app.core.config import JOB_ROLES_PATH

# Configure logging
logger = logging.getLogger(__name__)

STOPWORDS = {
    "a", "an", "and", "are", "as", "at", "be", "by", "for", "from", "has",
    "he", "in", "is", "it", "its", "of", "on", "that", "the", "to", "was",
    "were", "will", "with", "or", "this", "you", "your",
}

GENERIC_SKILL_TERMS = {
    "ability", "candidate", "collaboration", "communication", "degree",
    "engineer", "engineering", "environment", "experience", "framework",
    "frameworks", "knowledge", "preferred", "problem", "problems",
    "professional", "requirements", "responsibilities", "role", "skills",
    "software", "solutions", "strong", "team", "tools", "work",
}

NORMALIZATION_RULES = (
    (re.compile(r"(?<![a-z0-9])c\+\+(?![a-z0-9])", re.IGNORECASE), " cpp "),
    (re.compile(r"(?<![a-z0-9])c#(?![a-z0-9])", re.IGNORECASE), " csharp "),
    (re.compile(r"\bnode(?:\.|\s+)js\b", re.IGNORECASE), " nodejs "),
    (re.compile(r"\breact(?:\.|\s+)js\b", re.IGNORECASE), " reactjs "),
    (re.compile(r"\bnext(?:\.|\s+)js\b", re.IGNORECASE), " nextjs "),
    (re.compile(r"\bvue(?:\.|\s+)js\b", re.IGNORECASE), " vuejs "),
    (re.compile(r"\bexpress(?:\.|\s+)js\b", re.IGNORECASE), " expressjs "),
    (re.compile(r"\bangular(?:\.|\s+)js\b", re.IGNORECASE), " angularjs "),
    (re.compile(r"(?<![a-z0-9])\.net(?![a-z0-9])", re.IGNORECASE), " dotnet "),
)

SKILL_ALIAS_MAP = {
    "react": ["reactjs", "react js"],
    "angular": ["angularjs", "angular js"],
    "vue": ["vuejs", "vue js"],
    "nodejs": ["node", "node js"],
    "nextjs": ["next", "next js"],
    "express": ["expressjs", "express js"],
    "javascript": ["js"],
    "typescript": ["ts"],
    "rest api": ["rest", "restful api", "rest apis", "api"],
    "spring boot": ["springboot"],
    "ci cd": ["cicd", "continuous integration", "continuous delivery"],
    "ui ux": ["uiux", "ui", "ux"],
    "scikit learn": ["sklearn", "scikit"],
    "machine learning": ["ml"],
    "deep learning": ["dl"],
    "artificial intelligence": ["ai"],
    "shell scripting": ["shell script", "bash scripting", "bash"],
    "github actions": ["github action"],
    "tailwind css": ["tailwind"],
    "authentication": ["auth"],
    "jwt": ["json web token"],
    "jpa": ["java persistence api"],
    "nosql": ["no sql"],
    "microservices": ["microservice"],
    "data visualization": ["data visualisation"],
}

MAX_RESUME_TOKENS = 2800
MAX_JOB_DESCRIPTION_TOKENS = 1200
MAX_SKILLS = 40
SEMANTIC_CHUNK_SIZE = 180

_SEMANTIC_MODEL = None
_SEMANTIC_ENGINE = None


def _apply_normalization_rules(text: str) -> str:
    normalized = text.lower()

    for pattern, replacement in NORMALIZATION_RULES:
        normalized = pattern.sub(replacement, normalized)

    return normalized


def preprocess(text: str) -> str:
    text = _apply_normalization_rules(text)
    text = re.sub(r"[^a-z0-9\s]", " ", text)
    return re.sub(r"\s+", " ", text).strip()


def tokenize(text: str):
    return [
        token for token in preprocess(text).split()
        if token and token not in STOPWORDS
    ]


def _clean_extracted_text(text: str) -> str:
    normalized = unicodedata.normalize("NFKC", text or "")
    normalized = normalized.replace("\u00a0", " ")
    normalized = re.sub(r"([A-Za-z])-\s*\n\s*([A-Za-z])", r"\1\2", normalized)
    normalized = re.sub(r"(?<=\w)\s*\n\s*(?=\w)", " ", normalized)
    return re.sub(r"[ \t]+", " ", normalized).strip()


@lru_cache(maxsize=1)
def _load_role_bank():
    try:
        with open(JOB_ROLES_PATH, "r", encoding="utf-8") as file_handle:
            payload = json.load(file_handle)
    except Exception as exc:
        logger.warning("Could not load job role skill bank: %s", exc)
        return []

    return payload.get("job_roles", [])


def _role_skill_list(role: dict, job_desc: str = ""):
    skills = []

    for key in ("keywords", "soft_skills", "tools"):
        skills.extend(role.get(key, []))

    normalized_job = preprocess(job_desc)
    skills.extend(
        skill for skill in role.get("bonus_skills", [])
        if _count_skill_mentions(normalized_job, skill) > 0
    )

    return _dedupe_skills(skills)


def _infer_role_bank_skills(job_desc: str):
    normalized_job = preprocess(job_desc)
    scored_roles = []

    for role in _load_role_bank():
        score = 0
        role_name = role.get("role", "")
        role_terms = [term for term in tokenize(role_name) if term not in {"developer", "engineer"}]

        if preprocess(role_name) and preprocess(role_name) in normalized_job:
            score += 8

        score += sum(3 for term in role_terms if re.search(rf"(?<![a-z0-9]){re.escape(term)}(?![a-z0-9])", normalized_job))
        score += sum(1 for skill in role.get("keywords", []) if _count_skill_mentions(normalized_job, skill) > 0)

        if score > 0:
            scored_roles.append((score, role))

    if not scored_roles:
        return []

    scored_roles.sort(key=lambda item: item[0], reverse=True)
    best_score, best_role = scored_roles[0]

    if best_score < 3:
        return []

    return _role_skill_list(best_role, job_desc)


def _limit_text_by_tokens(text: str, max_tokens: int):
    compact_text = re.sub(r"\s+", " ", text).strip()

    if not compact_text:
        return "", 0, 0, False

    tokens = compact_text.split()
    total_tokens = len(tokens)

    if total_tokens <= max_tokens:
        return compact_text, total_tokens, total_tokens, False

    return " ".join(tokens[:max_tokens]), total_tokens, max_tokens, True


def _dedupe_skills(skills: list):
    unique_skills = []
    seen = set()

    for raw_skill in skills:
        cleaned_skill = re.sub(r"\s+", " ", raw_skill).strip(" ,;:-")
        normalized_skill = preprocess(cleaned_skill)

        if not cleaned_skill or not normalized_skill or normalized_skill in seen:
            continue

        seen.add(normalized_skill)
        unique_skills.append(cleaned_skill)

        if len(unique_skills) >= MAX_SKILLS:
            break

    return unique_skills


def _is_skill_candidate(candidate: str) -> bool:
    normalized_candidate = preprocess(candidate)

    if not normalized_candidate or normalized_candidate in GENERIC_SKILL_TERMS:
        return False

    parts = normalized_candidate.split()

    if not parts or len(parts) > 4:
        return False

    if all(part in STOPWORDS or part in GENERIC_SKILL_TERMS for part in parts):
        return False

    return len("".join(parts)) >= 2


def _split_skill_segment(segment: str):
    candidates = []

    sentence_chunks = re.split(r"\.(?=\s+[A-Z])", segment)

    for chunk in sentence_chunks:
        for piece in re.split(r",|;|\||\band\b", chunk, flags=re.IGNORECASE):
            cleaned_piece = re.sub(
                r"^(?:requirements?|required skills|preferred skills|skills?|tech stack|stack|technologies|tools?|strong|solid|working|hands on|proficiency|familiarity|experience with|experience in|knowledge of|expertise in|understanding of|using|with|in)\s*[:\-]?\s*",
                "",
                piece.strip(),
                flags=re.IGNORECASE,
            )
            cleaned_piece = re.sub(
                r"\s+(?:is|are)\s+(?:preferred|required|desired|helpful|a plus).*$",
                "",
                cleaned_piece,
                flags=re.IGNORECASE,
            )
            cleaned_piece = re.sub(
                r"\s+(?:preferred|required|desired|helpful|mandatory|bonus|plus)$",
                "",
                cleaned_piece,
                flags=re.IGNORECASE,
            )
            cleaned_piece = re.sub(r"\s+", " ", cleaned_piece).strip(" -:.")

            if _is_skill_candidate(cleaned_piece):
                candidates.append(cleaned_piece)

    return candidates


def _infer_skills_from_job_description(job_desc: str):
    inferred = []
    seen = set()
    context_patterns = (
        r"(?:required skills|requirements|must have|preferred skills|tech stack|stack|technologies|tools)\s*[:\-]\s*([^\n]+)",
        r"(?:experience with|proficient in|knowledge of|familiarity with|expertise in)\s+([^\n]+)",
    )

    for pattern in context_patterns:
        for match in re.finditer(pattern, job_desc, flags=re.IGNORECASE):
            for candidate in _split_skill_segment(match.group(1)):
                normalized_candidate = preprocess(candidate)
                if normalized_candidate not in seen:
                    seen.add(normalized_candidate)
                    inferred.append(candidate)

    for line in job_desc.splitlines():
        compact_line = line.strip(" -*\t")
        if not compact_line:
            continue

        if any(separator in compact_line for separator in [",", "|"]):
            for candidate in _split_skill_segment(compact_line):
                normalized_candidate = preprocess(candidate)
                if normalized_candidate not in seen:
                    seen.add(normalized_candidate)
                    inferred.append(candidate)

    if inferred:
        return _dedupe_skills(inferred)

    token_counts = Counter(
        token
        for token in tokenize(job_desc)
        if len(token) > 1 and token not in GENERIC_SKILL_TERMS
    )
    fallback_skills = []

    for token, _count in token_counts.most_common(MAX_SKILLS):
        if token in {"aws", "gcp", "sql", "api", "rest", "etl", "ml", "ai", "ci", "cd"}:
            fallback_skills.append(token.upper())
        else:
            fallback_skills.append(token)

    return _dedupe_skills(fallback_skills)


def _skill_aliases(skill: str):
    normalized_skill = preprocess(skill)

    if not normalized_skill:
        return []

    aliases = {normalized_skill}
    aliases.update(SKILL_ALIAS_MAP.get(normalized_skill, []))

    if normalized_skill.endswith("js") and len(normalized_skill) > 2:
        aliases.add(normalized_skill[:-2])

    if normalized_skill.endswith(" api"):
        aliases.add(normalized_skill[:-4])

    for alias in list(aliases):
        if " " not in alias and 2 <= len(alias) <= 12:
            aliases.add(" ".join(alias))

    return sorted(preprocess(alias) for alias in aliases if preprocess(alias))


@lru_cache(maxsize=512)
def _compile_skill_patterns(skill: str):
    patterns = []

    for alias in _skill_aliases(skill):
        parts = [re.escape(part) for part in alias.split()]

        if not parts:
            continue

        phrase_pattern = r"\s+".join(parts)
        patterns.append(re.compile(rf"(?<![a-z0-9]){phrase_pattern}(?![a-z0-9])"))

    return tuple(patterns)


def _count_skill_mentions(normalized_text: str, skill: str) -> int:
    patterns = _compile_skill_patterns(skill)
    return max((len(pattern.findall(normalized_text)) for pattern in patterns), default=0)


def extract_hard_keywords(text: str, keyword_list: list) -> set:
    normalized_text = preprocess(text)
    return {
        skill for skill in keyword_list
        if _count_skill_mentions(normalized_text, skill) > 0
    }


def _vector_cosine_similarity(left_text: str, right_text: str) -> float:
    left_counter = Counter(tokenize(left_text))
    right_counter = Counter(tokenize(right_text))
    vocabulary = set(left_counter) | set(right_counter)

    if not vocabulary:
        return 0.0

    dot_product = sum(left_counter[token] * right_counter[token] for token in vocabulary)
    left_norm = math.sqrt(sum(value * value for value in left_counter.values()))
    right_norm = math.sqrt(sum(value * value for value in right_counter.values()))

    if left_norm == 0 or right_norm == 0:
        return 0.0

    return dot_product / (left_norm * right_norm)


def _fallback_semantic_similarity(left_text: str, right_text: str) -> float:
    left_tokens = set(tokenize(left_text))
    right_tokens = set(tokenize(right_text))

    if not left_tokens or not right_tokens:
        return 0.0

    overlap = len(left_tokens & right_tokens) / len(left_tokens | right_tokens)
    return overlap


def _load_semantic_model():
    global _SEMANTIC_MODEL
    global _SEMANTIC_ENGINE

    if _SEMANTIC_ENGINE is not None:
        return _SEMANTIC_MODEL, _SEMANTIC_ENGINE

    try:
        from sentence_transformers import SentenceTransformer, util

        _SEMANTIC_MODEL = (SentenceTransformer("all-MiniLM-L6-v2"), util)
        _SEMANTIC_ENGINE = "sentence-transformers"
        return _SEMANTIC_MODEL, _SEMANTIC_ENGINE
    except Exception:
        _SEMANTIC_MODEL = None
        _SEMANTIC_ENGINE = "keyword-overlap"
        return None, _SEMANTIC_ENGINE


def _chunk_text(text: str):
    tokens = text.split()

    if not tokens:
        return []

    return [
        " ".join(tokens[index:index + SEMANTIC_CHUNK_SIZE])
        for index in range(0, len(tokens), SEMANTIC_CHUNK_SIZE)
    ]


def _semantic_similarity(left_text: str, right_text: str):
    model_bundle, engine = _load_semantic_model()

    if model_bundle is None:
        return _fallback_semantic_similarity(left_text, right_text), engine

    model, util = model_bundle
    left_chunks = _chunk_text(left_text)
    right_chunks = _chunk_text(right_text)

    if not left_chunks or not right_chunks:
        return 0.0, engine

    try:
        left_embeddings = model.encode(left_chunks, convert_to_tensor=True)
        right_embeddings = model.encode(right_chunks, convert_to_tensor=True)
        left_embedding = left_embeddings.mean(dim=0, keepdim=True)
        right_embedding = right_embeddings.mean(dim=0, keepdim=True)
        score = util.pytorch_cos_sim(left_embedding, right_embedding).item()
        return max(score, 0.0), engine
    except Exception:
        return _fallback_semantic_similarity(left_text, right_text), "keyword-overlap"


def _build_skill_importance_weights(job_desc: str, target_skills: list):
    normalized_job_desc = preprocess(job_desc)
    raw_weights = {}

    for skill in target_skills:
        raw_weights[skill] = 1 + min(_count_skill_mentions(normalized_job_desc, skill), 3)

    total_weight = sum(raw_weights.values()) or 1
    return {
        skill: weight / total_weight
        for skill, weight in raw_weights.items()
    }


def _calculate_skill_alignment_score(resume_text: str, target_skills: list, importance_weights: dict):
    normalized_resume = preprocess(resume_text)
    found_skills = []
    skill_mentions = {}
    weighted_score = 0.0

    for skill in target_skills:
        mentions = _count_skill_mentions(normalized_resume, skill)
        skill_mentions[skill] = mentions

        if mentions > 0:
            found_skills.append(skill)

        evidence_score = 0.0 if mentions == 0 else min(1.0, 0.82 + (mentions - 1) * 0.09)
        weighted_score += importance_weights.get(skill, 0.0) * evidence_score

    return weighted_score * 100, found_skills, skill_mentions


def _build_scoring_weights(skill_count: int, skill_source: str):
    if skill_count == 0:
        return {"skills": 0.0, "content": 0.55, "semantic": 0.45}

    if skill_count >= 10:
        skill_weight = 0.72
    elif skill_count >= 5:
        skill_weight = 0.68
    else:
        skill_weight = 0.62

    if skill_source == "job_description":
        skill_weight -= 0.05

    skill_weight = max(0.55, min(skill_weight, 0.75))
    remaining_weight = 1 - skill_weight

    return {
        "skills": round(skill_weight, 4),
        "content": round(remaining_weight * 0.6, 4),
        "semantic": round(remaining_weight * 0.4, 4),
    }


def _build_suggestion(
    skill_source: str,
    missing_skills: list,
    dynamic_skill_score: float,
    resume_truncated: bool,
    job_truncated: bool,
):
    notes = []

    if skill_source == "content_only":
        notes.append("Add required skills or a richer job description for stronger skill-based ATS scoring.")
    elif missing_skills:
        notes.append(
            "Add evidence for these missing skills or link them to real projects: "
            + ", ".join(missing_skills[:5])
            + "."
        )
    elif dynamic_skill_score < 75:
        notes.append("Matched skills are present, but the resume needs stronger proof through outcomes and project detail.")
    else:
        notes.append("Strong skill alignment. Tighten impact statements and measurable outcomes for a stronger ATS profile.")

    if resume_truncated or job_truncated:
        notes.append("Large inputs were safely shortened behind the scenes to keep ATS scoring stable.")

    return " ".join(notes)


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        from PyPDF2 import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [page.extract_text() or "" for page in reader.pages]
        text = _clean_extracted_text("\n".join(pages))
        logger.info(f"Extracted text from PDF: {len(text)} characters")
        return text
    except Exception as e:
        logger.error(f"Error extracting PDF: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Error parsing PDF file: {str(e)}")


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        import docx
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [paragraph.text for paragraph in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)

        text = _clean_extracted_text("\n".join(parts))
        logger.info(f"Extracted text from DOCX: {len(text)} characters")
        return text
    except Exception as e:
        logger.error(f"Error extracting DOCX: {str(e)}")
        raise HTTPException(status_code=400, detail=f"Error parsing DOCX file: {str(e)}")


async def parse_resume_file(upload: UploadFile) -> str:
    try:
        if not upload.filename:
            raise HTTPException(status_code=400, detail="File has no filename")
        
        file_bytes = await upload.read()
        
        if not file_bytes:
            raise HTTPException(status_code=400, detail="File is empty")
        
        filename = upload.filename.lower()
        logger.info(f"Parsing resume file: {upload.filename} ({len(file_bytes)} bytes)")

        if filename.endswith(".pdf"):
            return extract_text_from_pdf(file_bytes)
        elif filename.endswith(".docx"):
            return extract_text_from_docx(file_bytes)
        elif filename.endswith(".txt"):
            text = _clean_extracted_text(file_bytes.decode("utf-8", errors="ignore"))
            logger.info(f"Extracted text from TXT: {len(text)} characters")
            return text
        else:
            logger.warning(f"Unsupported file format: {filename}")
            raise HTTPException(
                status_code=400,
                detail="Unsupported resume format. Upload PDF, DOCX, or TXT.",
            )
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error parsing resume file: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error processing resume file")


def analyze_resume(resume_text: str, job_desc: str, required_skills: list):
    try:
        logger.info(f"Starting resume analysis: {len(resume_text)} chars, {len(required_skills)} required skills")
        
        target_skills = _dedupe_skills(required_skills)
        skill_source = "required_skills"

        if not target_skills:
            role_skills = _infer_role_bank_skills(job_desc)
            target_skills = role_skills or _infer_skills_from_job_description(job_desc)
            skill_source = "role_bank" if role_skills else "job_description" if target_skills else "content_only"

        normalized_resume_text, resume_token_count, analyzed_resume_tokens, resume_truncated = _limit_text_by_tokens(
            resume_text,
            MAX_RESUME_TOKENS,
        )
        normalized_job_desc, job_description_token_count, analyzed_job_description_tokens, job_truncated = _limit_text_by_tokens(
            job_desc,
            MAX_JOB_DESCRIPTION_TOKENS,
        )

        importance_weights = _build_skill_importance_weights(normalized_job_desc, target_skills)
        skill_alignment_score, found_skills, skill_mentions = _calculate_skill_alignment_score(
            resume_text,
            target_skills,
            importance_weights,
        )
        hard_skills_match = (len(found_skills) / len(target_skills) * 100) if target_skills else 0.0
        dynamic_skill_score = (
            0.65 * skill_alignment_score + 0.35 * hard_skills_match
            if target_skills else
            0.0
        )

        tfidf_score = _vector_cosine_similarity(normalized_resume_text, normalized_job_desc) * 100
        semantic_score_raw, semantic_engine = _semantic_similarity(normalized_resume_text, normalized_job_desc)
        semantic_score = semantic_score_raw * 100

        scoring_weights = _build_scoring_weights(len(target_skills), skill_source)
        final_score = (
            scoring_weights["skills"] * dynamic_skill_score +
            scoring_weights["content"] * tfidf_score +
            scoring_weights["semantic"] * semantic_score
        )
        missing_skills = sorted(list(set(target_skills) - set(found_skills)))
        
        logger.info(f"Resume analysis complete: {len(found_skills)}/{len(target_skills)} skills matched, overall score: {final_score:.2f}")

        return {
            "overall_match": round(final_score, 2),
            "hard_skills_match": round(hard_skills_match, 2),
            "skill_alignment_score": round(dynamic_skill_score, 2),
            "content_similarity": round(tfidf_score, 2),
            "semantic_similarity": round(semantic_score, 2),
            "target_skills": target_skills,
            "found_skills": sorted(found_skills),
            "missing_skills": missing_skills,
            "matched_skill_count": len(found_skills),
            "total_skill_count": len(target_skills),
            "skill_source": skill_source,
            "suggestion": _build_suggestion(
                skill_source,
                missing_skills,
                dynamic_skill_score,
                resume_truncated,
                job_truncated,
            ),
            "parsed_resume_length": len(resume_text),
            "resume_token_count": resume_token_count,
            "job_description_token_count": job_description_token_count,
            "analyzed_resume_tokens": analyzed_resume_tokens,
            "analyzed_job_description_tokens": analyzed_job_description_tokens,
            "resume_truncated": resume_truncated,
            "job_description_truncated": job_truncated,
            "scoring_weights": scoring_weights,
            "semantic_engine": semantic_engine,
            "skill_mentions": skill_mentions,
        }
    except Exception as e:
        logger.error(f"Error analyzing resume: {str(e)}", exc_info=True)
        raise HTTPException(status_code=500, detail="Error analyzing resume")
